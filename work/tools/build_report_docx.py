"""Build a visual bid-workbench result report (DOCX) from workbench files."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ROOT = Path(__file__).resolve().parent.parent
PROJECT = ROOT / "projects" / "N5100012026002044"
OUT = ROOT / "docs" / "招投标智能技能运行结果报告.docx"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK = RGBColor(0x1F, 0x4D, 0x78)
RED = RGBColor(0x9B, 0x1C, 0x1C)
GOLD = RGBColor(0x7A, 0x5A, 0x00)
INK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)

TAG_COLORS = {
    "[已确认]": RGBColor(0x1F, 0x3A, 0x5F),
    "[合理拟定]": RGBColor(0x2E, 0x74, 0xB5),
    "[待补充]": GOLD,
    "[待确认]": GOLD,
    "[禁止直接使用]": RED,
}

USABLE_DXA = 9360


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_geometry(table, widths_cm) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if index >= len(widths_cm):
                continue
            cell.width = Cm(widths_cm[index])
            cell.vertical_alignment = 1


def add_table(doc, headers, rows, widths_cm, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(text))
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = DARK
        set_cell_shading(cell, "E8EEF5")
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            text = "" if value is None else str(value)
            cell = cells[index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            for tag, color in TAG_COLORS.items():
                if tag in text:
                    run = paragraph.add_run(tag)
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = color
                    text = text.replace(tag, "")
            run = paragraph.add_run(text)
            run.font.size = Pt(font_size)
            run.font.color.rgb = INK
    set_table_geometry(table, widths_cm)
    return table


def add_tagged_paragraph(doc, text, size=10.5):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    remaining = text
    pattern = re.compile("(" + "|".join(re.escape(t) for t in TAG_COLORS) + ")")
    for part in pattern.split(remaining):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.font.size = Pt(size)
        if part in TAG_COLORS:
            run.font.color.rgb = TAG_COLORS[part]
            run.bold = True
        else:
            run.font.color.rgb = INK
    return paragraph


def read(name):
    return (PROJECT / name).read_text(encoding="utf-8")


def table_rows(text, prefix):
    rows = []
    for line in text.splitlines():
        if not line.startswith("| "):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if not cells or not cells[0].startswith(prefix):
            continue
        rows.append(cells)
    return rows


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    rpr = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rpr.set(qn("w:eastAsia"), "微软雅黑")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, DARK, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "微软雅黑"
        )
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.text = "招投标智能技能运行结果报告 · 四川省科智汇川人工智能科技有限责任公司"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GRAY

    footer = section.footer.paragraphs[0]
    footer.text = "N5100012026002044 · 技能生成辅助稿件，最终提交前须人工复核"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GRAY


def build():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    add_header_footer(section)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("招投标智能技能运行结果报告")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = DARK
    run.font.name = "Calibri"
    run.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run(
        "AI赋能铁道移动通信课程建设及智能体开发与应用课程资源建设（二次）"
    )
    run.font.size = Pt(12)
    run.font.color.rgb = BLUE

    add_table(
        doc,
        ["项目", "内容", "项目", "内容"],
        [
            ["项目编号", "N5100012026002044", "采购人", "四川铁道职业学院"],
            ["采购方式", "竞争性磋商", "最高限价", "650,000.00 元"],
            ["服务期", "合同签订之日起 300 日", "响应有效期", "不少于 90 天"],
            ["投标主体", "四川省科智汇川人工智能科技有限责任公司", "当前状态", "不可直接提交"],
        ],
        [3.0, 5.6, 3.0, 5.6],
        font_size=9,
    )

    doc.add_paragraph()
    callout = doc.add_paragraph()
    callout.paragraph_format.space_before = Pt(4)
    callout.paragraph_format.space_after = Pt(12)
    run = callout.add_run(
        "本报告由招投标智能技能自动生成。技能严格遵循原文优先、证据绑定、"
        "废标与评分分离三项约束；凡无企业证据支撑的内容一律标记，不予推断。"
    )
    run.font.size = Pt(10)
    run.font.color.rgb = DARK
    run.italic = True

    doc.add_heading("一、工作台运行概况", level=1)
    add_table(
        doc,
        ["工作台文件", "条目数", "产出角色", "作用"],
        [
            ["01_要求矩阵", "30", "投标助理", "招标要求 → 响应 → 证据映射"],
            ["02_废标项清单", "12", "投标助理", "无效响应风险条目"],
            ["03_评分点清单", "7", "投标助理", "分值结构与符号核对"],
            ["04_任务清单", "14", "投标助理", "可执行任务拆解"],
            ["10_证据索引", "15", "文档秘书", "证据编号与适用条款"],
            ["11_资料缺口清单", "19", "文档秘书", "缺失资料与影响范围"],
            ["30_提纲及初稿", "5", "售前助手", "提纲、方案、应答表初稿"],
            ["40_问题清单", "11", "文档审查员", "三级问题与回流目标"],
            ["41_附件签章检查表", "12", "文档审查员", "附件齐备与盖章位置"],
        ],
        [5.4, 2.0, 3.0, 6.8],
    )

    doc.add_heading("二、评分结构与符号核对", level=1)
    score_rows = [
        [r[0], r[2], r[3], r[4]]
        for r in table_rows(read("03_评分点清单.md"), "SCORE")
        if len(r) >= 5
    ]
    add_table(
        doc,
        ["评分项", "分值", "属性", "当前得分条件"],
        score_rows,
        [5.0, 1.8, 2.2, 8.2],
    )
    doc.add_paragraph()
    for line in [
        "[已确认] 符号计数核对：▲12、●10、█3，与第五章评分细则声明完全一致。",
        "[已确认] 关键技术条款：650,000 元限价、不接受联合体、90 天有效期、"
        "5% 履约保证金、300 日服务期、40%+60% 付款、30 分钟解密。",
        "[合理拟定] 主观分（实施 9 分 + 售后 6 分）可通过文字质量争取；"
        "客观分（技术 48 分、履约 22 分、业绩 5 分）必须依赖企业真实证据。",
    ]:
        add_tagged_paragraph(doc, line)

    doc.add_heading("三、废标与无效响应风险", level=1)
    void_rows = table_rows(read("02_废标项清单.md"), "VOID")
    add_table(
        doc,
        ["编号", "风险", "必须动作", "状态"],
        [[r[0], r[1], r[4], r[5]] for r in void_rows],
        [2.2, 4.6, 6.6, 2.8],
    )
    doc.add_paragraph()
    add_tagged_paragraph(
        doc,
        "[已确认] 评分项未满足通常按细则扣分；除非同时标注 ★，不得误判为废标项。"
        "技能已将两类条目强制分表，校验脚本会拦截混表。",
    )

    doc.add_page_break()
    doc.add_heading("四、资料缺口", level=1)
    gap_rows = table_rows(read("11_资料缺口清单.md"), "GAP")
    add_table(
        doc,
        ["编号", "缺少资料", "影响条款/评分点", "缺失后果"],
        [[r[0], r[1], r[2], r[3]] for r in gap_rows],
        [2.0, 5.4, 4.6, 4.2],
    )
    doc.add_paragraph()
    add_tagged_paragraph(
        doc,
        "[待补充] 当前最大风险不是方案文字，而是企业证据缺口。"
        "在 GAP-001、GAP-002、GAP-005、GAP-014 关闭之前，不应提交响应文件。",
    )

    doc.add_heading("五、审查结论与问题分布", level=1)
    problem_rows = table_rows(read("40_问题清单.md"), "Q-")
    counts = {"阻断": 0, "重大": 0, "一般": 0}
    for row in problem_rows:
        if row[1] in counts:
            counts[row[1]] += 1
    add_table(
        doc,
        ["问题级别", "数量", "含义", "处理要求"],
        [
            ["阻断", str(counts["阻断"]), "可能导致无效响应", "必须关闭后方可提交"],
            ["重大", str(counts["重大"]), "影响得分或响应完整性", "需补证据或修改内容"],
            ["一般", str(counts["一般"]), "格式与措辞", "定稿前统一处理"],
        ],
        [3.0, 2.0, 6.0, 6.2],
    )
    doc.add_paragraph()
    add_tagged_paragraph(
        doc,
        "[待人工] 当前阶段为 GATHER（资料补充）。11 个阻断与重大问题未关闭，"
        "项目状态判定为不可直接提交。",
    )

    doc.add_page_break()
    doc.add_heading("六、项目实施方案（评分细则 9 分项）", level=1)
    add_tagged_paragraph(
        doc,
        "[已确认] 本章覆盖评分细则要求的六个组成部分，"
        "对应 SCORE-005；缺一项扣 1.5 分，每处缺陷扣 0.75 分。",
    )
    plan = read("31_方案初稿/实施方案.md").splitlines()
    for line in plan:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("|"):
            continue
        else:
            add_tagged_paragraph(doc, stripped)

    doc.add_page_break()
    doc.add_heading("七、售后服务方案（评分细则 6 分项）", level=1)
    add_tagged_paragraph(
        doc,
        "[已确认] 本章覆盖评分细则要求的四个组成部分，对应 SCORE-006。",
    )
    after_sales = read("31_方案初稿/售后服务方案.md").splitlines()
    for line in after_sales:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("|"):
            continue
        else:
            add_tagged_paragraph(doc, stripped)

    doc.add_heading("八、技能约束与校验结果", level=1)
    add_table(
        doc,
        ["约束机制", "作用", "校验方式"],
        [
            ["原文可追溯", "每条要求带原文、章节、页码、符号、风险等级", "人工审查"],
            ["废标评分分离", "防止无效条款与得分条款混淆", "校验脚本拦截"],
            ["符号计数核对", "▲12 ●10 █3 与细则一致", "校验脚本"],
            ["内容标记强制", "方案正文必须带五级标记", "校验脚本"],
            ["证据绑定", "企业事实必须绑定证据编号", "证据索引"],
            ["问题分类回流", "按 ANALYZE/GATHER/DRAFT 分流", "审查员判定"],
            ["放行门槛", "阻断问题未关闭禁止进入提交", "总控状态机"],
        ],
        [4.0, 7.2, 5.0],
    )
    doc.add_paragraph()
    add_tagged_paragraph(
        doc,
        "[已确认] 校验命令：python skills/scripts/validate_workbench.py "
        "projects/N5100012026002044；输出 errors 为空方为合规。",
    )

    doc.add_heading("九、下一步行动", level=1)
    for text in [
        "[待补充] 补齐企业证明资料：营业执照、中小企业声明信息、软著证书、"
        "业绩合同、产品截图、三级等保备案、承诺函。",
        "[待确认] 企业决策报价与成本测算，确认不超过 650,000 元。",
        "[待确认] 确认人员配置名单与售后到场时限，方可写入实施方案定稿。",
        "[待人工] 完成数字证书与电子印章准备，落实 30 分钟解密值守安排。",
    ]:
        add_tagged_paragraph(doc, text)

    doc.save(str(OUT))
    print(f"written: {OUT}")


if __name__ == "__main__":
    sys.exit(build())
